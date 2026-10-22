"""동료 검증 결과의 일치도 계산.  사용: python agreement.py 동료검증_작성본.xlsx
(실험 폴더에서 실행: contrast.jsonl 필요)"""
import json, sys
from collections import Counter
from openpyxl import load_workbook
KEY_B = {"석공사(건물외부 공사)|조적공사": ["석공사(건물외부공사)"], "벽체, 천장 및 바닥의 공사|단열공사": ["단열공사"], "기둥·내력벽|기둥·내력벽": ["내력구조부"], "지하저수조(물탱크)공사|옥외급수ㆍ위생 관련 공사": ["저수조(물탱크)공사"], "·공기조화기기설비공사|난방·냉방·환기·공기조화설비공사": ["공기조화기기설비공사"], "옥외설비공사(담장,울타리,안내시설물 등)|잡공사": ["옥외설비공사"], "온돌공사|난방·냉방·환기, 공기조화설비공사": ["온돌공사(세대매립배관 포함)"], "타일공사(방수공사)|마감공사": ["타일공사"], "석공사/건물내부공사|마감공사": ["석공사(건물내부 공사)"], "옥외위생관련공사|옥외급수·위생관련공사": ["옥외위생(정화조) 관련 공사"], "옹벽공사|대지조성공사": ["옹벽공사(토목옹벽)"], "구조용 철골공사|철골공사": ["일반철골공사"], "블럭공사|조적공사": ["블록공사"], "보·바닥 및 지붕공사|보·바닥 및 지붕공사": ["내력구조부"], "벽체, 천장 및 바닥의 단열공사|": ["단열공사"]}
def kappa(pairs):
    n = len(pairs)
    if n == 0: return None
    po = sum(a == b for a, b in pairs) / n
    ca = Counter(a for a, _ in pairs); cb = Counter(b for _, b in pairs)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / (n * n)
    return (po - pe) / (1 - pe) if pe < 1 else 1.0
G = {json.loads(l)['case_id']: json.loads(l) for l in open('data/contrast.jsonl', encoding='utf-8')}
wb = load_workbook(sys.argv[1], data_only=True)
A = wb['A_대조사례']; lab, art, dis = [], [], []
for r in range(2, A.max_row + 1):
    cid, v, a, memo = A.cell(r, 2).value, A.cell(r, 6).value, A.cell(r, 7).value, A.cell(r, 8).value
    if not cid or not v: continue
    g = G[cid]; gl = g['label']; ga = '제' + g['cited_articles'][0]
    lab.append((gl, v)); art.append((ga, a))
    if gl != v or ga != a: dis.append((cid, g['title'], g['contrast_value'], f'정답 {gl}/{ga}', f'검증자 {v}/{a}', memo or ''))
B = wb['B_공종매핑']; bp, bdis = [], []
for r in range(2, B.max_row + 1):
    sub, maj, v = B.cell(r, 2).value, B.cell(r, 3).value or '', B.cell(r, 6).value
    if not sub or not v: continue
    gold = KEY_B.get(f'{sub}|{maj}', [])
    ok = v in gold; bp.append(ok)
    if not ok: bdis.append((sub, maj, f'연구자 {gold}', f'검증자 {v}', B.cell(r, 7).value or ''))
print(f'[시트 A] 입력 {len(lab)}건')
if lab:
    print(f'  하자 여부 일치율 {sum(a==b for a,b in lab)/len(lab):.1%}, Cohen κ = {kappa(lab):.3f}')
    print(f'  근거 조항 일치율 {sum(a==b for a,b in art)/len(art):.1%}')
print(f'[시트 B] 입력 {len(bp)}건, 일치율 {sum(bp)/len(bp):.1%}' if bp else '[시트 B] 입력 없음')
if dis or bdis: print('\n불일치 목록')
for d in dis: print('  A', *d, sep=' | ')
for d in bdis: print('  B', *d, sep=' | ')
